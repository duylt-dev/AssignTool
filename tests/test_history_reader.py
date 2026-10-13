from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest

from scrumbot.errors import InputError
from scrumbot.model import SprintLabel
from scrumbot.sources.history_reader import read_history
from tests.conftest import TECH_HEADERS, build_history_xlsx, write_history_csv


def _row(**values: object) -> dict:
    return {
        "Ngày lên đơn": "2026-09-22 09:36",
        "ID": "n_1",
        "Tên task": "1.0.1 - Add ads",
        "Loại": "Task",
        "Task cha": None,
        "Trạng thái": "done",
        "App Name": "Demo",
        "Người làm": "Ngọc Văn Linh",
        "Sprint": "Sprint 9 (9/21 - 10/2)",
        **values,
    }


def test_xlsx_reads_root_orders_of_every_month_sheet_and_skips_the_summary(tmp_path: Path) -> None:
    path = tmp_path / "tech.xlsx"
    build_history_xlsx(
        path,
        {
            "Tháng 9": [
                _row(
                    **{
                        "Người làm": "Ngọc Văn Linh, Trần Đức An",
                        "Sprint": "Sprint 18 (4/17 - 4/29), Sprint 19 (5/6 - 5/20)",
                        "Ngày lên đơn": "2026-04-16 10:00",
                        "Ưu tiên": "urgent",
                        "Due date": "2026-04-20 18:00",
                    }
                ),
                _row(**{"Tên task": "QA testing", "Task cha": "1.0.1 - Add ads"}),
                _row(**{"Loại": "Bug", "Tên task": "Bug 1"}),
            ],
            "Tháng 10": [_row(**{"Ngày lên đơn": datetime(2026, 10, 1, 9, 0), "Sprint": None, "App Name": None})],
        },
    )

    first, second = read_history(path)

    assert first.assignees == ("Ngọc Văn Linh", "Trần Đức An")
    assert first.sprint == "Sprint 18 (4/17 - 4/29), Sprint 19 (5/6 - 5/20)"
    assert first.sprints == (
        SprintLabel(18, date(2026, 4, 17), date(2026, 4, 29)),
        SprintLabel(19, date(2026, 5, 6), date(2026, 5, 20)),
    )
    assert (first.created, first.priority_raw, first.due, first.hours) == (
        date(2026, 4, 16),
        "urgent",
        date(2026, 4, 20),
        None,
    )
    assert (second.sprint, second.sprints, second.app, second.created) == ("", (), "", date(2026, 10, 1))


@pytest.mark.parametrize(("bom", "delimiter", "date_style"), [(True, ",", "iso"), (False, ";", "dmy")])
def test_csv_reads_supported_shapes_and_optional_hours(tmp_path: Path, bom: bool, delimiter: str, date_style: str):
    path = tmp_path / "tech.csv"
    write_history_csv(
        path,
        [_row(**{"Ngày lên đơn": "2026-10-06", "Ước lượng (h)": "16,5"}), _row(**{"Task cha": "Parent"})],
        bom=bom,
        delimiter=delimiter,
        date_style=date_style,
        headers=[*TECH_HEADERS, "Ước lượng (h)"],
    )

    (order,) = read_history(path)

    assert (order.created, order.hours, order.app) == (date(2026, 10, 6), 16.5, "Demo")


def test_missing_required_column_names_the_file_and_sheet(tmp_path: Path) -> None:
    csv_path = tmp_path / "tech.csv"
    write_history_csv(csv_path, [_row()], headers=[h for h in TECH_HEADERS if h != "Task cha"])
    with pytest.raises(InputError, match="tech.csv' thiếu cột 'Task cha'"):
        read_history(csv_path)

    xlsx_path = tmp_path / "tech.xlsx"
    build_history_xlsx(xlsx_path, {"Tháng 9": [_row()]})
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path)
    workbook["Tháng 9"].cell(1, TECH_HEADERS.index("Sprint") + 1).value = "Removed"
    workbook.save(xlsx_path)
    with pytest.raises(InputError, match="sheet 'Tháng 9' thiếu cột 'Sprint'"):
        read_history(xlsx_path)


def test_unreadable_or_empty_date_names_the_row(tmp_path: Path) -> None:
    csv_path = tmp_path / "tech.csv"
    write_history_csv(csv_path, [_row(**{"Ngày lên đơn": "hôm qua"})])
    with pytest.raises(InputError, match="dòng 2, cột 'Ngày lên đơn': không đọc được ngày"):
        read_history(csv_path)

    xlsx_path = tmp_path / "tech.xlsx"
    build_history_xlsx(xlsx_path, {"Tháng 9": [_row(), _row(**{"Ngày lên đơn": None})]})
    with pytest.raises(InputError, match="sheet 'Tháng 9', dòng 3, cột 'Ngày lên đơn' không được để trống"):
        read_history(xlsx_path)


def test_xlsx_without_a_data_sheet_is_an_input_error(tmp_path: Path) -> None:
    path = tmp_path / "tech.xlsx"
    build_history_xlsx(path, {})

    with pytest.raises(InputError, match="không có sheet nào có cột 'Ngày lên đơn'"):
        read_history(path)
