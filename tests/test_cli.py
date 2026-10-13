"""The command line against a real local HTTP server (assignment spec §2, §10, §11; bot spec §1, §6): input
discovery, exit codes, the report's sheets, the TM rules and the AI wiring. Nothing reaches the dev server."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scrumbot.actions.xlsx_report import report_name
from scrumbot.cli import main
from scrumbot.model import AiReply
from scrumbot.model_board import TECH_MANAGER_LIST
from scrumbot.understanding import ai_cli
from tests.api_builders import DEV_ROWS, LINH, item, orders, sprint, sprints
from tests.conftest import build_devs
from tests.http_server import FakeDevServer

DONE = item("old", "1.0.0", people=(LINH,), status="done", hours=12, created=datetime(2026, 9, 23, 9))
EMPTY = item("new", "1.0.1 - Add ads", status="ready to progress", created=datetime(2026, 10, 7, 9))
RULES = "# Luật\n\n| Loại | Dev | App | Ghi chú |\n|---|---|---|---|\n| Chỉ định | MinhNM | Demo | thử |\n"


@pytest.fixture
def server(monkeypatch: pytest.MonkeyPatch) -> Iterator[FakeDevServer]:
    running = FakeDevServer(
        {
            "sprints": sprints(
                sprint("S1", "Sprint 1 (9/22 - 10/3)", date(2026, 9, 22), date(2026, 10, 3), archived=True),
                sprint("S2", "Sprint 2 (10/6 - 10/17)", date(2026, 10, 6), date(2026, 10, 17)),
            ),
            "don": {"S1": orders(DONE), "S2": orders(EMPTY), TECH_MANAGER_LIST: orders(DONE, EMPTY)},
            "lich-su": {"ok": True, "ids": [], "khong_thay": [], "doi_trang_thai": [], "comment": [], "nguoi": []},
        }
    )
    monkeypatch.setenv("SCRUMBOT_API_BASE", running.base)
    yield running
    running.close()


@pytest.fixture
def source(tmp_path: Path) -> Path:
    folder = tmp_path / "input"
    folder.mkdir()
    build_devs(folder / "devs.xlsx", DEV_ROWS)
    return folder


def _run(source: Path, tmp_path: Path, *flags: str) -> int:
    paths = ["--input", str(source), "--state", str(tmp_path / "state"), "--out", str(tmp_path / "out")]
    return main(["run", *paths, *flags])


def _report(tmp_path: Path) -> Path:
    (path,) = (tmp_path / "out").glob("assignment_sprint-2_*.xlsx")
    return path


def _assignment(tmp_path: Path) -> dict[str, object]:
    header, row = list(load_workbook(_report(tmp_path))["Phân công"].iter_rows(values_only=True))
    return dict(zip(header, row, strict=True))


def test_off_run_writes_the_six_sheets_and_prints_the_proposal(
    server: FakeDevServer, source: Path, tmp_path: Path, capsys
) -> None:
    assert _run(source, tmp_path, "--ai", "off", "--today", "2026-10-08") == 0

    workbook = load_workbook(_report(tmp_path))
    assert workbook.sheetnames == ["Tóm tắt", "Phân công", "Rủi ro", "Thay đổi", "Cảnh báo", "Tải dev"]
    row = _assignment(tmp_path)
    assert (row["Độ lớn"], row["Giờ ước tính"], row["Mã PIC"], row["Nguồn"]) == (
        "M (mặc định)",
        16,
        "LinhNV",
        "Tool đề xuất",
    )
    assert str(row["Lý do"]).startswith("Làm version gần nhất của Demo (Sprint 1 (9/22 - 10/3): 1.0.0)")
    printed = capsys.readouterr().out
    assert "Ngày chạy: 08/10/2026 (--today)" in printed and "Độ lớn: tắt AI (mọi đơn M)" in printed
    assert "Luật TM: không có" in printed
    assert "STT 1 · Demo 1.0.1 - Add ads [M] → LinhNV (dự phòng: MinhNM, AnTD)" in printed
    assert not list((tmp_path / "out").glob("*.tmp"))


def test_ai_reply_sizes_the_order_without_names_and_with_the_history_hours(
    server: FakeDevServer, source: Path, tmp_path: Path, monkeypatch, capsys
) -> None:
    calls = []

    def fake(tool, prompt, *, model, timeout):
        calls.append((tool, model, timeout, LINH in prompt, "- Demo · 1.0.0 · 12h" in prompt))
        return AiReply(tool, '[{"stt": 1, "size": "L", "reason": "Thêm ads"}]', None)

    monkeypatch.setattr(ai_cli, "run_ai", fake)

    assert _run(source, tmp_path, "--ai-model", "m1", "--ai-timeout", "9", "--today", "2026-10-08") == 0

    assert calls == [("gemini", "m1", 9, False, True)]
    row = _assignment(tmp_path)
    assert (row["Độ lớn"], row["Giờ ước tính"], row["Lý do độ lớn"]) == ("L", 32, "Thêm ads")
    assert "Độ lớn: AI (gemini, 1/1 đơn)" in capsys.readouterr().out


def test_ai_failure_falls_back_to_m_and_still_writes_the_report(
    server: FakeDevServer, source: Path, tmp_path: Path, monkeypatch, capsys
) -> None:
    monkeypatch.setattr(ai_cli, "run_ai", lambda tool, prompt, **_: AiReply(tool, None, "mã thoát 41"))

    assert _run(source, tmp_path, "--today", "2026-10-08") == 0

    printed = capsys.readouterr().out
    assert "Độ lớn: mặc định M (mã thoát 41)" in printed and "WARN SIZE-AI-FAILED" in printed
    findings = list(load_workbook(_report(tmp_path))["Cảnh báo"].iter_rows(values_only=True))
    assert findings[1][:2] == ("WARN", "SIZE-AI-FAILED")


def test_the_md_file_in_the_input_folder_is_the_rules_file_and_its_pin_reaches_the_report(
    server: FakeDevServer, source: Path, tmp_path: Path, capsys
) -> None:
    (source / "luat.md").write_text(RULES.replace("\n", "\r\n"), encoding="utf-8-sig")

    assert _run(source, tmp_path, "--ai", "off", "--today", "2026-10-08") == 0

    printed = capsys.readouterr().out
    assert "Luật TM: file luat.md (1 luật, 1 có hiệu lực ngày 08/10/2026)" in printed
    assert "STT 1 · Demo 1.0.1 - Add ads [M] → MinhNM " in printed
    assert str(_assignment(tmp_path)["Lý do"]).startswith("Luật TM dòng 5: chỉ định MinhNM cho Demo (thử)")


@pytest.mark.parametrize(("today", "pic"), [("2026-10-14", "MinhNM"), ("2026-10-08", "LinhNV")])
def test_today_replaces_the_clock_for_the_rules_day(
    server: FakeDevServer, source: Path, tmp_path: Path, capsys, today, pic
) -> None:
    """A pin dated 12/10–31/10 acts on 14/10 but not on 08/10 (A34)."""
    dated = (
        "| Loại | Dev | App | Từ ngày | Đến ngày |\n|---|---|---|---|---|\n"
        "| Chỉ định | MinhNM | Demo | 12/10/2026 | 31/10/2026 |\n"
    )
    (source / "luat.md").write_text(dated, encoding="utf-8")

    assert _run(source, tmp_path, "--ai", "off", "--today", today) == 0

    assert f"STT 1 · Demo 1.0.1 - Add ads [M] → {pic} " in capsys.readouterr().out


@pytest.mark.parametrize(
    ("files", "flags", "message"),
    [
        (["~$devs.xlsx"], [], "Không tìm thấy file devs.xlsx. Chỉ định bằng --devs."),
        (["devs.xlsx", "a.md", "b.md"], [], "Có nhiều file luật .md: a.md, b.md. Chỉ định một file bằng --rules."),
        (["devs.xlsx", "a.md"], [], "không có bảng luật (bảng có cột Loại và Dev)."),
        (["devs.xlsx"], ["--rules", "missing.md"], "Không có file 'missing.md' (--rules)."),
        (["devs.xlsx", "a.md:rules"], [], "dòng 5: không có dev mã 'Nobody' trong file dev."),
        (["devs.xlsx"], ["--today", "31/02/2026"], "--today '31/02/2026': không đọc được ngày"),
    ],
)
def test_input_errors_exit_2_with_a_message(tmp_path: Path, capsys, files, flags, message) -> None:
    source = tmp_path / "input"
    source.mkdir()
    for spec in files:
        name, _, kind = spec.partition(":")
        if name == "devs.xlsx":
            build_devs(source / name, DEV_ROWS)
        elif kind == "rules":
            (source / name).write_text(RULES.replace("MinhNM", "Nobody"), encoding="utf-8")
        else:
            (source / name).write_bytes(b"")

    assert _run(source, tmp_path, "--ai", "off", *flags) == 2

    assert message in capsys.readouterr().err
    assert not (tmp_path / "out").exists()


def test_missing_input_directory_exits_2(tmp_path: Path, capsys) -> None:
    assert _run(tmp_path / "nope", tmp_path, "--ai", "off") == 2
    assert "Không có thư mục input" in capsys.readouterr().err


def test_report_name_uses_a_windows_safe_slug() -> None:
    now = datetime(2026, 10, 6, 14, 30, 0)
    assert report_name("Sprint 10", now) == "assignment_sprint-10_20261006-143000.xlsx"
    assert report_name("Sprint: 1/2", now) == "assignment_sprint_-1_2_20261006-143000.xlsx"
