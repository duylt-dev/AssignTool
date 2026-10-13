from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scrumbot.errors import InputError
from scrumbot.sources.sprint_reader import read_workbook
from tests.conftest import SPRINT_HEADERS, build_sprint_workbook


def _row(stt: int, **values: object) -> dict:
    return {
        "STT": stt,
        "Task": "Build feature",
        "Loại": "Task",
        "App Name.": "App",
        "Trạng thái": "In progress",
        "Người làm": "Ngo\u0323c Va\u0306n Linh, Bùi Văn Tiến",
        "Ưu tiên": "high",
        "Hạn": datetime(2026, 10, 9),
        "Ngày tạo": datetime(2026, 10, 6),
        **values,
    }


def test_reads_rows_dates_and_stops_before_formula_footer(tmp_path: Path) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 9": [_row(1)], "Sprint 10": [_row(2), _row(3, Loại="Subtask", Task="Child")]})

    target, previous, previous_name = read_workbook(path)

    assert target.name == "Sprint 10"
    assert target.start == date(2026, 10, 6)
    assert target.end == date(2026, 10, 17)
    assert len(target.rows) == 2
    assert target.rows[0].assignees[0] == "Ngọc Văn Linh"
    assert target.rows[1].title == "Child"
    assert all("COUNTIF" not in str(cell) for row in target.rows for cell in row.cells)
    assert previous is not None and previous.name == "Sprint 9"
    assert previous_name == "Sprint 9"


def test_selects_override_and_reports_invalid_sheet(tmp_path: Path) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 9": [_row(1)], "Sprint 10": [_row(2)]})

    target, _, _ = read_workbook(path, "Sprint 9")

    assert target.name == "Sprint 9"
    with pytest.raises(InputError, match="Sprint 9, Sprint 10"):
        read_workbook(path, "Sprint 8")


def test_missing_previous_and_absent_row_three_dates(tmp_path: Path) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 10": [_row(1)]}, {"Sprint 10": ("unknown", "unknown")})

    target, previous, previous_name = read_workbook(path)

    assert (target.start, target.end) == (None, None)
    assert previous is None
    assert previous_name == "Sprint 9"


@pytest.mark.parametrize(
    "header", ["STT", "Task", "Loại", "App Name.", "Trạng thái", "Người làm", "Ưu tiên", "Hạn", "Ngày tạo"]
)
def test_target_requires_each_header(tmp_path: Path, header: str) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 10": [_row(1)]})
    workbook = load_workbook(path)
    worksheet = workbook["Sprint 10"]
    worksheet.cell(5, SPRINT_HEADERS.index(header) + 1).value = "Removed"
    workbook.save(path)

    with pytest.raises(InputError, match="thiếu cột|tiêu đề STT"):
        read_workbook(path)


@pytest.mark.parametrize("header", ["Loại", "Trạng thái", "Người làm"])
def test_previous_requires_its_headers(tmp_path: Path, header: str) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 9": [_row(1)], "Sprint 10": [_row(2)]})
    workbook = load_workbook(path)
    worksheet = workbook["Sprint 9"]
    worksheet.cell(5, SPRINT_HEADERS.index(header) + 1).value = "Removed"
    workbook.save(path)

    with pytest.raises(InputError, match="Sprint 9.*thiếu cột"):
        read_workbook(path)


def test_bad_target_date_names_excel_row_and_column(tmp_path: Path) -> None:
    path = tmp_path / "sprint.xlsx"
    build_sprint_workbook(path, {"Sprint 10": [_row(1, Hạn="not a date")]})

    with pytest.raises(InputError, match="dòng 6, cột 'Hạn'"):
        read_workbook(path)


def test_read_sheet_history_reads_only_older_sheets_and_their_hours(tmp_path) -> None:
    from datetime import datetime

    from scrumbot.sources.sprint_reader import read_sheet_history
    from tests.conftest import build_sprint_workbook

    task = {"Loại": "Task", "App Name.": "Demo", "Người làm": "Dev A", "Ngày tạo": datetime(2026, 9, 1)}
    path = tmp_path / "export.xlsx"
    build_sprint_workbook(
        path,
        {
            "Sprint 10": [{**task, "STT": 1, "Task": "new", "Ước lượng (h)": 99}],
            "Sprint 9": [
                {**task, "STT": 1, "Task": "v2", "Ước lượng (h)": 12},
                {**task, "STT": 2, "Task": "sub", "Loại": "Subtask", "Ước lượng (h)": 3},
                {**task, "STT": 3, "Task": "no app", "App Name.": ""},
            ],
            "Sprint 8": [{**task, "STT": 1, "Task": "v1", "Ước lượng (h)": "1,5"}],
        },
    )

    orders, names = read_sheet_history(path, before=10)

    assert names == ("Sprint 8", "Sprint 9")
    assert [(o.sprint, o.title, o.hours) for o in orders] == [("Sprint 8", "v1", 1.5), ("Sprint 9", "v2", 12.0)]
