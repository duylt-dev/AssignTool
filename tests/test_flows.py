"""The run end to end against a real local HTTP server (bot spec §1–§6): two days with the memory in between.
Nothing reaches the dev server."""

from __future__ import annotations

import copy
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scrumbot.cli import main
from scrumbot.flows.common import Options
from scrumbot.flows.run_api import run_api
from scrumbot.model_board import TECH_MANAGER_LIST
from tests.api_builders import AN, DEV_ROWS, LINH, MINH, TIEN, history, item, orders, sprint, sprints
from tests.conftest import build_devs
from tests.http_server import FakeDevServer

TUESDAY = datetime(2026, 10, 13, 9, 0)
ROBLOX = item("a", "1.0.4", people=(AN,), app="Roblox", status="done", created=datetime(2026, 9, 20, 9))
PS = item(
    "b", "1.1.1", people=(MINH, LINH), app="PS Remote", status="need to be clean", created=datetime(2026, 9, 19, 9)
)
EMPTY = item("c", "1.0.5", app="Roblox", status="ready to progress", created=datetime(2026, 10, 1, 9))
DEMO = item("e", "2.0.0", people=(LINH,), app="Demo", created=datetime(2026, 10, 2, 9))
GAME = item("f", "1.2.0", people=(TIEN,), app="Game", created=datetime(2026, 10, 2, 9))
BACKLOG = item("g", "9.9.9", app="Other", status="all issue", created=datetime(2026, 10, 10, 9))


def _data() -> dict[str, object]:
    return {
        "sprints": sprints(
            sprint("S9", "Sprint 9 (9/21 - 10/2)", date(2026, 9, 21), date(2026, 10, 2), archived=True, total=2),
            sprint("S10", "Sprint 10 (10/6 - 10/17)", date(2026, 10, 6), date(2026, 10, 17), total=3),
            sprint("T10", "[Test sprint] Sprint 10 (10/6 - 10/17)", date(2026, 10, 6), date(2026, 10, 17), total=3),
        ),
        "don": {
            "S9": orders(ROBLOX, PS),
            "S10": orders(EMPTY, DEMO, GAME),
            "T10": orders(EMPTY, DEMO, GAME),
            TECH_MANAGER_LIST: orders(ROBLOX, PS, EMPTY, DEMO, GAME, BACKLOG),
        },
        "lich-su": history(("e", "ready to progress", "in progress", datetime(2026, 10, 7, 9))),
    }


@pytest.fixture
def server(monkeypatch: pytest.MonkeyPatch) -> Iterator[FakeDevServer]:
    running = FakeDevServer(_data())
    monkeypatch.setenv("SCRUMBOT_API_BASE", running.base)
    yield running
    running.close()


@pytest.fixture
def options(tmp_path: Path) -> Options:
    (tmp_path / "input").mkdir()
    build_devs(tmp_path / "input" / "devs.xlsx", DEV_ROWS)
    return Options(
        input_dir=tmp_path / "input",
        state_dir=tmp_path / "state",
        out_dir=tmp_path / "out",
        ai="off",
        env_file=tmp_path / "absent.env",
    )


def _sheet(path: Path, name: str) -> list[tuple]:
    return list(load_workbook(path)[name].iter_rows(values_only=True))


def test_a_manual_run_proposes_keeps_the_proposal_and_reports_the_changes_of_the_next_day(
    server: FakeDevServer, options: Options
) -> None:
    lines: list[str] = []
    first = run_api(options, lines.append, TUESDAY)

    assert "STT 1 · Roblox 1.0.5 [M] → AnTD (dự phòng: MinhNM, TiếnBV)" in lines
    assert any(line.startswith("Nguồn: API 127.0.0.1 (dữ liệu mới nhất 10/10/2026 09:00)") for line in lines)
    summary = dict(_sheet(first, "Tóm tắt")[1:])
    assert summary["Đơn trống chưa vào sprint"] == "1 đơn ở list Tech manager, 1 đơn tạo trong 30 ngày gần đây"
    assert summary["Thay đổi"] == "Chưa có lần chạy trước để so"
    header = _sheet(first, "Phân công")[0]
    assert header[:11] == (
        "STT",
        "ID",
        "Task",
        "App Name",
        "Trạng thái",
        "Người làm",
        "Ưu tiên",
        "Hạn",
        "Ngày tạo",
        "Bug mở",
        "Bug tổng",
    )
    assert "/api/doc/don" in server.requests and server.requests.count("/api/doc/sprints") == 1

    again: list[str] = []
    run_api(options, again.append, datetime(2026, 10, 13, 10, 0))
    assert "STT 1 · Roblox 1.0.5 [M] → AnTD (giữ đề xuất) (dự phòng: MinhNM, TiếnBV)" in again

    data = copy.deepcopy(_data())
    data["don"]["S10"] = orders({**EMPTY, "_people": [AN]}, {**DEMO, "uu_tien": "urgent"}, GAME)  # type: ignore[index]
    server.data = data
    next_day = run_api(options, [].append, datetime(2026, 10, 14, 9, 0))

    changes = _sheet(next_day, "Thay đổi")
    assert changes[1][0] == "So với dữ liệu lưu lúc 13/10/2026 10:00"
    assert ("Đổi ưu tiên", "STT 2 · 2.0.0", "trống", "urgent") in changes
    assert ("TM đã giao", "STT 1 · 1.0.5", "AnTD", "AnTD (khớp)") in changes
    risks = {row[1]: row for row in _sheet(next_day, "Rủi ro")[1:]}
    assert risks["ORDER-CHANGED"][:5] == (
        "Trung bình",
        "ORDER-CHANGED",
        "Thay đổi",
        "Đơn STT 2",
        "Đơn STT 2 đổi ưu tiên trống → urgent",
    )


def test_the_cli_refuses_the_production_server(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys) -> None:
    (tmp_path / "input").mkdir()
    build_devs(tmp_path / "input" / "devs.xlsx", DEV_ROWS)
    monkeypatch.setenv("SCRUMBOT_API_BASE", "https://circle.datlt.site")

    code = main(["run", "--input", str(tmp_path / "input"), "--state", str(tmp_path / "state"), "--ai", "off"])

    assert code == 2
    assert "trỏ tới production (circle.datlt.site); bot chỉ đọc server dev" in capsys.readouterr().err
