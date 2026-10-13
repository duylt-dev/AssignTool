"""Acceptance on the real sample: the tables of assignment spec §12, through `run --source file`, with --ai off and
the run day pinned to Sprint 10's first day (06/10/2026), so the result never depends on the clock."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from scrumbot.cli import main
from scrumbot.sources.sprint_reader import read_workbook
from tests.conftest import build_devs, write_history_csv
from tests.sample import DEVS, SAMPLE, SKIP_REASON, TECH, TECH_SKIP_REASON, history_rows

# PS Remote 1.1.2 is urgent with no `Hạn` (bot spec §3, D53); the only finding of runs A and E.
DUE_MISSING = [("WARN", "DUE-MISSING", "Đơn STT 7")]

pytestmark = [pytest.mark.real_input, pytest.mark.skipif(not SAMPLE.is_file(), reason=SKIP_REASON)]


def _run(tmp_path: Path, workbook: Path, *, csv: bool = True, history: Path | None = None) -> Path:
    """csv=True: history = sheet Sprint 9 as a Tech manager CSV (runs A, B). csv=False: no history file, so the older
    sheets (runs C, D). history: that file (run E)."""
    source = tmp_path / "input"
    source.mkdir()
    devs = source / "devs.xlsx"
    build_devs(devs, DEVS)
    if csv and history is None:
        _, previous, _ = read_workbook(SAMPLE)
        write_history_csv(source / "history.csv", history_rows(previous))
    out = tmp_path / "out"
    flags = ["--history", str(history)] if history else []
    code = main(
        ["run", "--source", "file", "--today", "2026-10-06", "--input", str(source), "--sprint", str(workbook)]
        + ["--devs", str(devs), "--out", str(out), "--ai", "off", *flags]
    )
    assert code == 0
    (report,) = out.glob("assignment_sprint-10_*.xlsx")
    return report


def _table(report: Path, sheet: str) -> list[dict]:
    workbook = load_workbook(report)
    rows = list(workbook[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=False)) for row in rows[1:]]


def _devs(report: Path) -> dict[str, tuple[int, int, str]]:
    return {
        r["Mã"]: (r["Việc dở sprint trước"], r["Tổng giờ ước tính"], r["App đã làm"]) for r in _table(report, "Tải dev")
    }


def test_run_a_keeps_the_five_preassigned_orders_and_proposes_minh_for_roblox(tmp_path: Path) -> None:
    report = _run(tmp_path, SAMPLE)

    rows = {r["STT"]: r for r in _table(report, "Phân công")}
    assert sorted(rows) == [1, 7, 10, 14, 18, 24]
    roblox = rows[10]
    assert (roblox["Mã PIC"], roblox["Nguồn"], roblox["Dự phòng 1"], roblox["Dự phòng 2"]) == (
        "MinhNM",
        "Tool đề xuất",
        "LinhNV",
        "TiếnBV",
    )
    assert roblox["Độ lớn"] == "M (mặc định)" and roblox["Giờ ước tính"] == 16
    assert {stt for stt, r in rows.items() if r["Nguồn"] == "Đã giao sẵn"} == {1, 7, 14, 18, 24}
    assert [(r["Mức"], r["Mã"], r["Vị trí"]) for r in _table(report, "Cảnh báo")] == DUE_MISSING
    assert _devs(report) == {
        "LinhNV": (1, 32, "AR Dog, metal detector, PS Remote, Santa prank, SpeakerCleaner"),
        "MinhNM": (0, 32, "PS Remote, Speak2"),
        "AnTD": (0, 32, "AR Dog, Game Emulator"),
        "TiếnBV": (2, 16, "Game mode 2"),
    }


def _cleared(tmp_path: Path) -> Path:
    """The sample with "Người làm" emptied on every Sprint 10 order (the backtests, runs B and D)."""
    cleared = tmp_path / "cleared.xlsx"
    workbook = load_workbook(SAMPLE)
    sheet = workbook["Sprint 10"]
    headers = [cell.value for cell in sheet[5]]
    kind, people = headers.index("Loại") + 1, headers.index("Người làm") + 1
    for row in range(6, sheet.max_row + 1):
        if sheet.cell(row, kind).value == "Task":
            sheet.cell(row, people).value = None
    workbook.save(cleared)
    return cleared


def test_run_b_backtest_without_assignees_matches_the_spec_table(tmp_path: Path) -> None:
    report = _run(tmp_path, _cleared(tmp_path))

    rows = _table(report, "Phân công")
    assert {r["STT"]: r["Mã PIC"] for r in rows} == {
        7: "MinhNM",
        24: "AnTD",
        14: "LinhNV",
        10: "MinhNM",
        1: "LinhNV",
        18: "TiếnBV",
    }
    assert all(r["Nguồn"] == "Tool đề xuất" for r in rows)
    assert {code: hours for code, (_, hours, _) in _devs(report).items()} == {
        "LinhNV": 32,
        "MinhNM": 32,
        "AnTD": 16,
        "TiếnBV": 16,
    }


def test_run_c_without_csv_reads_sprints_6_to_9_and_proposes_an_for_roblox(tmp_path: Path) -> None:
    report = _run(tmp_path, SAMPLE, csv=False)

    roblox = {r["STT"]: r for r in _table(report, "Phân công")}[10]
    assert (roblox["Mã PIC"], roblox["Dự phòng 1"], roblox["Dự phòng 2"]) == ("AnTD", "MinhNM", "TiếnBV")
    assert roblox["Lý do"].startswith("Làm version gần nhất của Roblox skin (Sprint 8: ")
    assert [(r["Mức"], r["Mã"], r["Vị trí"]) for r in _table(report, "Cảnh báo")] == [
        *DUE_MISSING,
        ("INFO", "HISTORY-FROM-WORKBOOK", "Sprint 6, Sprint 7, Sprint 8, Sprint 9"),
    ]
    assert {code: hours for code, (_, hours, _) in _devs(report).items()} == {
        "LinhNV": 32,
        "MinhNM": 16,
        "AnTD": 48,
        "TiếnBV": 16,
    }


def test_run_d_backtest_without_csv_matches_the_spec_table(tmp_path: Path) -> None:
    report = _run(tmp_path, _cleared(tmp_path), csv=False)

    assert {r["STT"]: r["Mã PIC"] for r in _table(report, "Phân công")} == {
        7: "MinhNM",
        24: "AnTD",
        14: "AnTD",
        10: "LinhNV",
        1: "MinhNM",
        18: "LinhNV",
    }
    assert {code: hours for code, (_, hours, _) in _devs(report).items()} == {
        "LinhNV": 32,
        "MinhNM": 32,
        "AnTD": 32,
        "TiếnBV": 0,
    }


@pytest.mark.skipif(not TECH.is_file(), reason=TECH_SKIP_REASON)
def test_run_e_tech_manager_history_proposes_an_for_roblox(tmp_path: Path) -> None:
    report = _run(tmp_path, SAMPLE, history=TECH)

    roblox = {r["STT"]: r for r in _table(report, "Phân công")}[10]
    assert (roblox["Mã PIC"], roblox["Dự phòng 1"], roblox["Dự phòng 2"]) == ("AnTD", "MinhNM", "TiếnBV")
    assert roblox["Lý do"].startswith("Làm version gần nhất của Roblox skin (Sprint 8 (9/4 - 9/18): 1.0.1 - Add ads)")
    assert [(r["Mức"], r["Mã"], r["Vị trí"]) for r in _table(report, "Cảnh báo")] == DUE_MISSING
    assert {code: hours for code, (_, hours, _) in _devs(report).items()} == {
        "LinhNV": 32,
        "MinhNM": 16,
        "AnTD": 48,
        "TiếnBV": 16,
    }
