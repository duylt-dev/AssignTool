"""Read one assignment report and score it against a case's dap-an.json (confirm.md A37, A40)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

CLASSES = ("S", "M", "L", "XL")


@dataclass
class Row:
    stt: int
    size: str
    default: bool
    pic: str
    source: str
    reason: str
    backups: tuple[str, ...]


@dataclass
class Report:
    rows: dict[int, Row]
    findings: list[tuple[str, str, str]]  # (code, where, message)


@dataclass
class Score:
    tm: tuple[int, int] = (0, 0)
    spec_pic: tuple[int, int] = (0, 0)
    spec_backups: tuple[int, int] = (0, 0)
    kept: tuple[int, int] = (0, 0)
    size_exact: tuple[int, int] = (0, 0)
    size_near: tuple[int, int] = (0, 0)
    size_default: int = 0
    found: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)
    extra: list[str] = field(default_factory=list)
    misses: list[str] = field(default_factory=list)


def read_report(path: Path) -> Report:
    workbook = load_workbook(path, read_only=True)
    sheet = list(workbook["Phân công"].iter_rows(values_only=True))
    index = {name: i for i, name in enumerate(sheet[0])}
    rows = {}
    for values in sheet[1:]:
        size = str(values[index["Độ lớn"]] or "")
        backups = tuple(str(values[index[c]]) for c in ("Dự phòng 1", "Dự phòng 2") if values[index[c]])
        stt = int(values[index["STT"]])
        rows[stt] = Row(
            stt=stt,
            size=size.split(" ")[0],
            default="(mặc định)" in size,
            pic=str(values[index["Mã PIC"]] or ""),
            source=str(values[index["Nguồn"]] or ""),
            reason=str(values[index["Lý do"]] or ""),
            backups=backups,
        )
    findings = [
        (str(code), str(where or ""), str(message or ""))
        for _, code, where, message in list(workbook["Cảnh báo"].iter_rows(values_only=True))[1:]
        if code
    ]
    return Report(rows, findings)


def _finding_key(code: str, stt: int | None) -> str:
    return f"{code} (STT {stt})" if stt else code


def score(report: Report, answer: dict, *, ai_run: bool) -> Score:
    result = Score()
    tm = spec = backups = kept = exact = near = total_tm = total_spec = total_kept = total_size = 0
    for order in answer["orders"]:
        row = report.rows.get(order["stt"])
        if row is None:  # a skipped row (TYPE-UNKNOWN) is not in the report
            continue
        if order["pre_assigned"]:
            total_kept += 1
            kept += row.source == "Đã giao sẵn"
            continue
        if order.get("tm"):
            total_tm += 1
            if row.pic in order["tm"]:
                tm += 1
            else:
                result.misses.append(f"STT {row.stt} {order['app']}: tool {row.pic}, TM {', '.join(order['tm'])}")
        if order.get("spec_pic") and not ai_run:
            total_spec += 1
            spec += row.pic == order["spec_pic"]
            backups += list(row.backups) == order["spec_backups"]
            if row.pic != order["spec_pic"] or list(row.backups) != order["spec_backups"]:
                result.misses.append(
                    f"STT {row.stt}: tool {row.pic} ({', '.join(row.backups)}), "
                    f"spec {order['spec_pic']} ({', '.join(order['spec_backups'])})"
                )
    for order in answer["orders"]:
        row = report.rows.get(order["stt"])
        if row is None or not ai_run:
            continue
        result.size_default += row.default
        if order.get("size_true") and row.size in CLASSES:
            total_size += 1
            gap = abs(CLASSES.index(row.size) - CLASSES.index(order["size_true"]))
            exact += gap == 0
            near += gap <= 1
    result.tm, result.spec_pic, result.spec_backups = (tm, total_tm), (spec, total_spec), (backups, total_spec)
    result.kept, result.size_exact, result.size_near = (kept, total_kept), (exact, total_size), (near, total_size)
    actual = {_finding_key(code, _stt(where, message)) for code, where, message in report.findings}
    actual = {key for key in actual if not key.startswith("SIZE-AI")}  # AI findings are reported separately
    expected = {_finding_key(f["code"], f["stt"]) for f in answer["findings"]}
    result.found = sorted(expected & actual)
    result.missing = sorted(expected - actual)
    result.extra = sorted(actual - expected)
    return result


def _stt(where: str, message: str) -> int | None:
    match = re.search(r"STT (\d+)", f"{where} {message}")
    return int(match[1]) if match else None
