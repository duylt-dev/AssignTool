"""Build the five evaluation cases in eval/ (confirm.md A36–A39). Run from the repo root:

    uv run python eval/build_cases.py

Each case gets `eval/<case>/input/` (what the tool reads: sprint workbook, devs.xlsx, rules .md) and, beside it,
`dap-an.md` (for people) and `dap-an.json` (for run_eval.py). The answers stay outside `input/`: an `.md` there would
be discovered as the rules file (D29). Only files under `eval/case-*/` are written or replaced.
"""

from __future__ import annotations

import json
import re
import shutil
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from scenarios import HISTORY, SYNTHETIC

ROOT = Path(__file__).resolve().parent.parent
EVAL = ROOT / "eval"
SOURCE = ROOT / "input" / "Maintain_update_Sprint_10_va_9.xlsx"
HEADERS = (
    "STT", "Nhóm trạng thái", "Task", "Loại", "Task cha", "App Name.", "Trạng thái", "Người làm", "Ưu tiên", "Hạn",
    "Ước lượng (h)", "Người tạo", "Điểm", "Git Pull Request", "Ngày tạo", "Mô tả",
)  # fmt: skip
POOL = (
    ("Ngọc Văn Linh", "LinhNV", "middle+"),
    ("Nguyễn Nhật Minh", "MinhNM", "middle"),
    ("Trần Đức An", "AnTD", "junior+"),
    ("Bùi Văn Tiến", "TiếnBV", "junior"),
)
CODES = {unicodedata.normalize("NFC", name).casefold(): code for name, code, _ in POOL}
CLASSES = ("S", "M", "L", "XL")


def size_class(hours: float) -> str:
    """Real hours to a size class (A40)."""
    return "S" if hours <= 8 else "M" if hours <= 24 else "L" if hours <= 48 else "XL"


def pool_codes(people: object) -> list[str]:
    names = (unicodedata.normalize("NFC", part).strip().casefold() for part in str(people or "").split(","))
    return [CODES[name] for name in names if name in CODES]


def write_sheet(workbook: Workbook, sheet: tuple[str, str, str, str], rows: list[dict], app_header: str) -> None:
    """One sheet in the export's layout (spec §3): title rows 1–3, header on row 5, a formula footer."""
    name, label, start, end = sheet
    headers = [app_header if h == "App Name." else h for h in HEADERS]
    worksheet = workbook.create_sheet(name)
    worksheet.append([label])
    worksheet.append(["Product Development / Maintain update"])
    worksheet.append([f"Thời gian: {start} → {end}   ·   Dữ liệu đánh giá (eval/)"])
    worksheet.append([])
    worksheet.append(headers)
    for row in rows:
        values = dict(row)
        if values.get("Loại") == "Subtask" and not str(values["Task"]).startswith(" "):
            values["Task"] = f"    └ {values['Task']}"
        worksheet.append([values.get(h) for h in HEADERS])
    worksheet.append(["Tổng (chỉ tính task gốc, tránh cộng trùng subtask)", '=COUNTIF(A:A,">0")'])
    worksheet.append([])
    worksheet.append(["Ghi chú: file sinh bởi eval/build_cases.py."])


def write_devs(path: Path, leave: dict[str, str], extra: list[tuple]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Tên", "Mã", "Level", "Nhận việc", "Ngày nghỉ"])
    for name, code, level in POOL:
        worksheet.append([name, code, level, "Y", leave.get(name)])
    for row in extra:
        worksheet.append(list(row))
    workbook.save(path)


def fresh(slug: str) -> Path:
    case = EVAL / slug
    if (case / "input").exists():
        shutil.rmtree(case / "input")
    (case / "input").mkdir(parents=True)
    return case


def read_source_sheet(book, name: str) -> tuple[tuple[str, str, str, str], list[dict]]:
    """One sheet of the sample export: (name, label, start, end) and its data rows as dicts."""
    source = list(book[name].iter_rows(values_only=True))
    header_row = next(i for i, row in enumerate(source[:10]) if row[0] == "STT")
    start, end = re.search(r"(\d{4}-\d{2}-\d{2})\s*→\s*(\d{4}-\d{2}-\d{2})", source[2][0]).groups()
    rows = []
    for values in source[header_row + 1 :]:
        if not isinstance(values[0], int):
            break
        rows.append(dict(zip(source[header_row], values, strict=False)))
    return (name, source[0][0], start, end), rows


def real_case(slug: str, sheet_name: str, title: str) -> dict:
    """A past sprint replayed (A38): root orders only, pool PICs and hours removed; older sheets kept for their hours,
    the previous one with every status set to done (carry-over 0)."""
    book = load_workbook(SOURCE, data_only=True)
    sheet, source_rows = read_source_sheet(book, sheet_name)
    start = sheet[2]
    number = int(sheet_name.split()[1])
    rows, orders = [], []
    for row in source_rows:
        if row["Loại"] != "Task":
            continue
        tm, hours = pool_codes(row["Người làm"]), row["Ước lượng (h)"]
        pre = bool(row["Người làm"]) and not tm
        row.update({"Ước lượng (h)": None, "Điểm": None, "Git Pull Request": None})
        row.update({"Trạng thái": "ready to progress", "Nhóm trạng thái": "ready to progress"})
        if tm:
            row["Người làm"] = None
        rows.append(row)
        hours = hours if isinstance(hours, int | float) and hours > 0 else None
        orders.append(
            {
                "stt": row["STT"],
                "app": row["App Name."],
                "title": str(row["Task"]).strip(),
                "priority": row["Ưu tiên"],
                "pre_assigned": pre,
                "tm": tm or None,
                "hours": hours,
                "size_true": size_class(hours) if hours else None,
            }
        )
    case = fresh(slug)
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_sheet(workbook, sheet, rows, "App Name.")
    for older in sorted((n for n in book.sheetnames if re.fullmatch(r"Sprint \d+", n)), key=lambda n: -int(n[7:])):
        if int(older[7:]) >= number:
            continue
        older_sheet, older_rows = read_source_sheet(book, older)
        if int(older[7:]) == number - 1:
            older_rows = [{**r, "Trạng thái": "done", "Nhóm trạng thái": "done"} for r in older_rows]
        write_sheet(workbook, older_sheet, older_rows, "App Name.")
    workbook.save(case / "input" / f"Maintain_{sheet_name.replace(' ', '_')}_dung_lai.xlsx")
    write_devs(case / "input" / "devs.xlsx", {}, [])
    due_before = [o["stt"] for o, r in zip(orders, rows, strict=True) if _due_before(r)]
    due_missing = [o["stt"] for o, r in zip(orders, rows, strict=True) if _due_missing(r)]
    return {
        "slug": slug,
        "title": title,
        "kind": "real",
        "args": ["--today", start, "--history", HISTORY],
        "orders": orders,
        "findings": [
            *({"code": "DUE-BEFORE-CREATED", "stt": s} for s in due_before),
            *({"code": "DUE-MISSING", "stt": s} for s in due_missing),
        ],
        "watch": [],
    }  # fmt: skip


def _due_missing(row: dict) -> bool:
    """D53, bot spec §3: an urgent or high order without `Hạn`."""
    return str(row["Ưu tiên"] or "").strip().casefold() in ("urgent", "high") and not row["Hạn"]


def _due_before(row: dict) -> bool:
    due, created = row["Hạn"], row["Ngày tạo"]
    return isinstance(due, datetime | date) and isinstance(created, datetime | date) and due < created


def synthetic_case(scenario: dict) -> dict:
    case = fresh(scenario["slug"])
    workbook = Workbook()
    workbook.remove(workbook.active)
    app_header = scenario.get("app_header", "App Name.")
    write_sheet(workbook, scenario["sprint"], scenario["rows"], app_header)
    if scenario.get("prev"):
        write_sheet(workbook, scenario["prev"], scenario["prev_rows"], app_header)
    workbook.save(case / "input" / "Maintain_Sprint_11.xlsx")
    write_devs(case / "input" / "devs.xlsx", scenario["devs"], scenario.get("extra_devs", []))
    if scenario.get("rules"):
        (case / "input" / "luat.md").write_text(scenario["rules"], encoding="utf-8")
    orders = []
    for row in scenario["rows"]:
        if row["Loại"] == "Subtask":
            continue
        pic, backups = scenario["expect"].get(row["STT"], (None, ()))
        orders.append(
            {
                "stt": row["STT"],
                "app": row["App Name."],
                "title": row["Task"],
                "priority": row["Ưu tiên"],
                "pre_assigned": bool(row["Người làm"]),
                "tm": None,
                "spec_pic": pic,
                "spec_backups": list(backups),
                "size_target": scenario["size_target"].get(row["STT"]),
                "type": row["Loại"],
            }
        )
    return {
        "slug": scenario["slug"],
        "title": scenario["title"],
        "kind": "synthetic",
        "args": ["--today", scenario["today"], "--history", HISTORY],
        "orders": orders,
        "findings": [{"code": code, "stt": stt} for code, stt in scenario["findings"]],
        "watch": scenario["watch"],
    }


def write_answers(answer: dict) -> None:
    case = EVAL / answer["slug"]
    (case / "dap-an.json").write_text(json.dumps(answer, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    command = " ".join(["uv run scrumbot run --source file --input", f"eval/{answer['slug']}/input", *answer["args"]])
    lines = [
        f"# {answer['slug']}: {answer['title']}",
        "",
        "Sinh bởi `eval/build_cases.py`; sửa ở đó, không sửa tay.",
        "",
    ]
    lines += [f"Chạy: `{command}`", ""]
    if answer["kind"] == "real":
        lines += ["Đáp án: dev TM đã giao (khớp khi PIC là một trong số đó) và giờ thật (A37, A40).", ""]
        lines += ["| STT | App | Đơn | Ưu tiên | TM giao | Giờ thật | Hạng thật |", "|---|---|---|---|---|---|---|"]
        for o in answer["orders"]:
            tm = ", ".join(o["tm"] or []) or ("giao sẵn (ngoài danh sách)" if o["pre_assigned"] else "")
            lines.append(
                f"| {o['stt']} | {o['app']} | {o['title']} | {o['priority'] or ''} | {tm} | "
                f"{o['hours'] or ''} | {o['size_true'] or ''} |"
            )
    else:
        lines += ["Đáp án theo spec, tính tay cho lần `--ai off` (A37). Cột `TM điền` để trống cho TM (O10).", ""]
        lines += [
            "| STT | App | Đơn | Ưu tiên | Giao sẵn | PIC theo spec | Dự phòng theo spec | Hạng nhắm tới | TM điền |",
            "|---|---|---|---|---|---|---|---|---|",
        ]
        for o in answer["orders"]:
            lines.append(
                f"| {o['stt']} | {o['app'] or ''} | {o['title']} | {o['priority'] or ''} | "
                f"{'✓' if o['pre_assigned'] else ''} | {o['spec_pic'] or ''} | {', '.join(o['spec_backups'])} | "
                f"{o['size_target'] or ''} | |"
            )
    expected = [f"`{f['code']}`" + (f" (STT {f['stt']})" if f["stt"] else "") for f in answer["findings"]]
    lines += ["", "Cảnh báo theo spec: " + (", ".join(expected) or "không có")]
    if answer["watch"]:
        lines += ["", "Cần quan sát (không chấm):", "", *(f"- {item}" for item in answer["watch"])]
    (case / "dap-an.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    answers = [
        real_case("case-1-sprint-9", "Sprint 9", "Sprint 9 dựng lại"),
        real_case("case-2-sprint-8", "Sprint 8", "Sprint 8 dựng lại"),
        *(synthetic_case(scenario) for scenario in SYNTHETIC),
    ]
    for answer in answers:
        write_answers(answer)
        print(f"{answer['slug']}: {len(answer['orders'])} đơn")


if __name__ == "__main__":
    main()
