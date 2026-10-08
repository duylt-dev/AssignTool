"""Write the Excel reports: the assignment (spec §10) and the backtest (spec §13). The only module that writes
openpyxl workbooks."""

from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from assigntool.model import Assignment, BacktestSprint, DevSummary, Finding, Level, Result
from assigntool.rules.backtest import percent

_EXTRA_HEADERS = (
    "Độ lớn",
    "Giờ ước tính",
    "Lý do độ lớn",
    "PIC",
    "Mã PIC",
    "Nguồn",
    "Lý do",
    "Dự phòng 1",
    "Dự phòng 2",
)
_FINDING_HEADERS = ("Mức", "Mã", "Vị trí", "Nội dung")
_DEV_HEADERS = (
    "Mã",
    "Tên",
    "Level",
    "Việc dở sprint trước",
    "Ngày nghỉ",
    "Đơn giao sẵn",
    "Đơn tool đề xuất",
    "Tổng đơn",
    "Tổng giờ ước tính",
    "Danh sách đơn",
    "App đã làm",
)
_SPRINT_HEADERS = (
    "Sprint",
    "Bắt đầu",
    "Kết thúc",
    "Đơn so sánh",
    "Khớp",
    "Tỉ lệ khớp",
    "Đơn mang sang",
    "Đơn theo luật TM",
)
_ORDER_HEADERS = (
    "Sprint",
    "Ngày lên đơn",
    "App",
    "Đơn",
    "Ưu tiên",
    "TM giao",
    "Tool giao",
    "Khớp",
    "Dự phòng 1",
    "Dự phòng 2",
    "Lý do của tool",
    "Luật TM",
)
_LEVEL_NAMES = {
    Level.JUNIOR: "junior",
    Level.JUNIOR_PLUS: "junior+",
    Level.MIDDLE: "middle",
    Level.MIDDLE_PLUS: "middle+",
    Level.SENIOR: "senior",
}
_MAX_WIDTH = 60


def report_name(sheet_name: str, now: datetime) -> str:
    """`Sprint 10` at 2026-10-06 14:30:00 → `assignment_sprint-10_20261006-143000.xlsx` (best practices §5)."""
    slug = re.sub(r"[^a-z0-9._-]", "_", re.sub(r"\s+", "-", sheet_name.strip().lower()))
    return f"assignment_{slug}_{now:%Y%m%d-%H%M%S}.xlsx"


def write_report(result: Result, headers: tuple[str, ...], out_dir: Path, now: datetime) -> Path:
    """Write the three sheets to `<name>.tmp`, then rename it (best practices §6)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / report_name(result.sheet_name, now)
    workbook = Workbook()
    _assignments_sheet(workbook.active, result.assignments, headers)
    _findings_sheet(workbook.create_sheet("Cảnh báo"), result.findings)
    _devs_sheet(workbook.create_sheet("Tải dev"), result.devs)
    temporary = path.with_name(path.name + ".tmp")
    workbook.save(temporary)
    os.replace(temporary, path)
    return path


def write_backtest(sprints: tuple[BacktestSprint, ...], out_dir: Path, now: datetime) -> Path:
    """`backtest_<YYYYMMDD-HHMMSS>.xlsx` with sheets `Theo sprint` and `Theo đơn`, via `.tmp` (spec §13)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"backtest_{now:%Y%m%d-%H%M%S}.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Theo sprint"
    summary.append(list(_SPRINT_HEADERS))
    for sprint in sprints:
        label = sprint.label
        matched = sum(order.matched for order in sprint.orders)
        summary.append(
            [label.name, label.start, label.end, *_counts(matched, len(sprint.orders)), sprint.kept, _ruled(sprint)]
        )
    matched = sum(order.matched for sprint in sprints for order in sprint.orders)
    total = sum(len(sprint.orders) for sprint in sprints)
    kept = sum(sprint.kept for sprint in sprints)
    summary.append(["Tổng", None, None, *_counts(matched, total), kept, sum(_ruled(sprint) for sprint in sprints)])
    _finish(summary)
    _date_format(summary, 2, 3)
    detail = workbook.create_sheet("Theo đơn")
    detail.append(list(_ORDER_HEADERS))
    for sprint in sprints:
        for item in sprint.orders:
            backups = (*item.assignment.backups, "", "")
            detail.append(
                [
                    sprint.label.name,
                    item.order.created,
                    item.order.app,
                    item.order.title,
                    item.order.priority_raw,
                    ", ".join(item.tm_codes),
                    item.assignment.pic_codes[0],
                    "✓" if item.matched else "✗",
                    backups[0],
                    backups[1],
                    item.assignment.reason,
                    ", ".join(f"dòng {line}" for line in item.assignment.rule_lines),
                ]
            )
    _finish(detail)
    _date_format(detail, 2, 2)
    temporary = path.with_name(path.name + ".tmp")
    workbook.save(temporary)
    os.replace(temporary, path)
    return path


def _ruled(sprint: BacktestSprint) -> int:
    """Compared orders whose choice a TM rule shaped (spec §13)."""
    return sum(bool(order.assignment.rule_lines) for order in sprint.orders)


def _counts(matched: int, total: int) -> list[object]:
    return [total, matched, f"{percent(matched, total)}%" if total else ""]


def _date_format(sheet: Worksheet, first: int, last: int) -> None:
    for row in sheet.iter_rows(min_row=2, min_col=first, max_col=last):
        for cell in row:
            cell.number_format = "DD/MM/YYYY"


def _assignments_sheet(sheet: Worksheet, assignments: tuple[Assignment, ...], headers: tuple[str, ...]) -> None:
    sheet.title = "Phân công"
    sheet.append([*headers, *_EXTRA_HEADERS])
    for item in assignments:
        cells = list(item.order.row.cells[: len(headers)])
        cells += [None] * (len(headers) - len(cells))
        size = item.size
        backups = (*item.backups, "", "")
        sheet.append(
            [
                *cells,
                size.size.name + (" (mặc định)" if size.is_default else ""),
                size.size.value,
                size.reason,
                ", ".join(item.pic_names),
                ", ".join(item.pic_codes),
                item.source.value,
                item.reason,
                backups[0],
                backups[1],
            ]
        )
    _finish(sheet)


def _findings_sheet(sheet: Worksheet, findings: tuple[Finding, ...]) -> None:
    sheet.append(list(_FINDING_HEADERS))
    if not findings:
        sheet.append(["Không có cảnh báo"])
    for finding in findings:
        sheet.append([finding.severity.value, finding.code, finding.where, finding.message])
    _finish(sheet)


def _devs_sheet(sheet: Worksheet, devs: tuple[DevSummary, ...]) -> None:
    sheet.append(list(_DEV_HEADERS))
    for item in devs:
        orders = "Nghỉ cả sprint" if item.on_leave_all_sprint else "\n".join(item.orders)
        sheet.append(
            [
                item.dev.code,
                item.dev.name,
                _LEVEL_NAMES[item.dev.level],
                item.carry_over,
                item.leave_days,
                item.preassigned,
                item.proposed,
                item.preassigned + item.proposed,
                item.hours,
                orders,
                ", ".join(item.known_apps),
            ]
        )
    _finish(sheet)


def _finish(sheet: Worksheet) -> None:
    """Bold, frozen header; wrapped text; widths from the longest line in each column, capped."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).splitlines() or [""])
                widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), _MAX_WIDTH)
