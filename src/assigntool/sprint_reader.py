"""Read sprint workbooks into value objects (spec §3)."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from assigntool.errors import InputError
from assigntool.model import HistoryOrder, SprintRow, SprintSheet
from assigntool.normalize import clean, header_key, parse_date, split_people, text_key

_SPRINT = re.compile(r"^Sprint\s+(\d+)$")
_DATES = re.compile(r"(\d{4}-\d{2}-\d{2})\s*→\s*(\d{4}-\d{2}-\d{2})")
_TARGET_HEADERS = ("stt", "task", "loại", "app name", "trạng thái", "người làm", "ưu tiên", "hạn", "ngày tạo")
_PREVIOUS_HEADERS = ("loại", "trạng thái", "người làm")
_HISTORY_HEADERS = ("loại", "app name", "người làm", "ngày tạo")


def read_workbook(path: Path, sheet: str | None = None) -> tuple[SprintSheet, SprintSheet | None, str]:
    """Read the target sprint, its preceding sprint, and the expected preceding name."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as error:
        raise InputError(f"Không mở được file '{path}': {error}") from error
    try:
        sprints = _sprint_sheets(workbook.sheetnames)
        if not sprints:
            raise InputError(f"File '{path}' không có sheet Sprint hợp lệ.")
        target_name, target_number = _select_target(path, sprints, sheet)
        previous_name = f"Sprint {target_number - 1}"
        target = _read_sheet(path, workbook[target_name], target_number, _TARGET_HEADERS)
        previous = None
        for name, number in sprints:
            if number == target_number - 1:
                previous = _read_sheet(path, workbook[name], number, _PREVIOUS_HEADERS)
                break
        return target, previous, previous_name
    finally:
        workbook.close()


def read_sheet_history(path: Path, before: int) -> tuple[tuple[HistoryOrder, ...], tuple[str, ...]]:
    """History from the workbook's sprint sheets numbered below `before`, used when there is no CSV (spec §5).

    Returns the orders ("Loại" = Task, with an app, a name and a creation date) and the sheet names read, oldest first.
    """
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as error:
        raise InputError(f"Không mở được file '{path}': {error}") from error
    try:
        orders: list[HistoryOrder] = []
        names: list[str] = []
        for name, number in sorted(_sprint_sheets(workbook.sheetnames), key=lambda item: item[1]):
            if number >= before:
                continue
            sheet = _read_sheet(path, workbook[name], number, _HISTORY_HEADERS)
            names.append(sheet.name)
            orders.extend(
                HistoryOrder(sheet.name, row.app, row.title, row.assignees, row.hours, row.created)
                for row in sheet.rows
                if text_key(row.kind) == "task" and row.app and row.assignees and row.created is not None
            )
        return tuple(orders), tuple(names)
    finally:
        workbook.close()


def _sprint_sheets(names: list[str]) -> list[tuple[str, int]]:
    return [(name, int(match[1])) for name in names if (match := _SPRINT.fullmatch(name.strip()))]


def _select_target(path: Path, sprints: list[tuple[str, int]], requested: str | None) -> tuple[str, int]:
    if requested is not None:
        for name, number in sprints:
            if name == requested:
                return name, number
        available = ", ".join(name for name, _ in sprints)
        raise InputError(f"File '{path}' không có sheet '{requested}'. Các sheet Sprint: {available}.")
    return max(sprints, key=lambda item: item[1])


def _read_sheet(path: Path, worksheet: object, number: int, required: tuple[str, ...]) -> SprintSheet:
    header_row, headers = _headers(path, worksheet)
    columns = {header_key(value): index for index, value in enumerate(headers) if header_key(value)}
    for key in required:
        if key not in columns:
            raise InputError(f"File '{path}', sheet '{worksheet.title}' thiếu cột '{_display_header(key)}'.")
    dates = _sheet_dates(worksheet)
    rows = tuple(_read_rows(path, worksheet, header_row, headers, columns))
    return SprintSheet(worksheet.title, number, *dates, tuple(clean(value) for value in headers), rows)


def _headers(path: Path, worksheet: object) -> tuple[int, tuple[object, ...]]:
    for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row and header_key(row[0]) == "stt":
            return row_number, tuple(row)
    raise InputError(f"File '{path}', sheet '{worksheet.title}' không tìm thấy hàng tiêu đề STT.")


def _sheet_dates(worksheet: object) -> tuple[date | None, date | None]:
    row = next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True), ())
    match = _DATES.search(" ".join(clean(value) for value in row))
    if not match:
        return None, None
    return date.fromisoformat(match[1]), date.fromisoformat(match[2])


def _read_rows(path: Path, worksheet: object, header_row: int, headers: tuple[object, ...], columns: dict[str, int]):
    for excel_row, raw in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True), start=header_row + 1
    ):
        cells = tuple(raw) + (None,) * max(0, len(headers) - len(raw))
        if not _is_stt(cells[0] if cells else None):
            return
        yield SprintRow(
            int(cells[0]),
            clean(_title(_cell(cells, columns, "task"))),
            clean(_cell(cells, columns, "loại")),
            clean(_cell(cells, columns, "app name")),
            split_people(_cell(cells, columns, "người làm")),
            clean(_cell(cells, columns, "trạng thái")),
            clean(_cell(cells, columns, "ưu tiên")),
            _date_cell(path, worksheet.title, excel_row, "hạn", cells, columns),
            _date_cell(path, worksheet.title, excel_row, "ngày tạo", cells, columns),
            clean(_cell(cells, columns, "mô tả")),
            cells,
            _hours(_cell(cells, columns, "ước lượng (h)")),
        )


def _is_stt(value: object) -> bool:
    return (isinstance(value, int) and not isinstance(value, bool)) or (isinstance(value, float) and value.is_integer())


def _hours(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(clean(value).replace(",", ".")) if clean(value) else None
    except ValueError:
        return None


def _title(value: object) -> object:
    if not isinstance(value, str):
        return value
    return value.lstrip().lstrip("└").lstrip()


def _cell(cells: tuple[object, ...], columns: dict[str, int], key: str) -> object:
    index = columns.get(key)
    return cells[index] if index is not None and index < len(cells) else None


def _date_cell(
    path: Path, sheet: str, excel_row: int, key: str, cells: tuple[object, ...], columns: dict[str, int]
) -> date | None:
    value = _cell(cells, columns, key)
    try:
        return parse_date(value)
    except ValueError as error:
        raise InputError(
            f"File '{path}', sheet '{sheet}', dòng {excel_row}, cột '{_display_header(key)}': {error}"
        ) from error


def _display_header(key: str) -> str:
    return {"app name": "App Name", "ngày tạo": "Ngày tạo", "người làm": "Người làm"}.get(key, key.title())
