"""Read devs.xlsx into Dev values (spec §4)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from scrumbot.core.errors import InputError
from scrumbot.core.model import Dev
from scrumbot.core.normalize import clean, header_key, parse_date, parse_level, text_key

_REQUIRED = ("tên", "mã", "level", "nhận việc")


def read_devs(path: Path) -> tuple[Dev, ...]:
    """Read all non-empty developer rows from the first workbook sheet."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as error:
        raise InputError(f"Không mở được file '{path}': {error}") from error
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        headers = tuple(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        columns = {header_key(value): index for index, value in enumerate(headers) if header_key(value)}
        for key in _REQUIRED:
            if key not in columns:
                raise InputError(f"File '{path}', sheet '{worksheet.title}' thiếu cột '{key.title()}'.")
        devs: list[Dev] = []
        names: set[str] = set()
        codes: set[str] = set()
        for excel_row, raw in enumerate(
            worksheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
        ):
            cells = tuple(raw)
            name, code, level_value = (_cell(cells, columns, key) for key in ("tên", "mã", "level"))
            if not any(clean(value) for value in (name, code, level_value)):
                continue
            level = parse_level(level_value)
            if level is None:
                raise InputError(
                    f"File '{path}', dòng {excel_row}, cột 'Level' không hợp lệ; "
                    "dùng junior, junior+, middle, middle+, senior."
                )
            active_key = text_key(_cell(cells, columns, "nhận việc"))
            if active_key not in {"y", "n"}:
                raise InputError(f"File '{path}', dòng {excel_row}, cột 'Nhận việc' phải là Y hoặc N.")
            clean_name, clean_code = clean(name), clean(code)
            if text_key(clean_name) in names:
                raise InputError(f"File '{path}', dòng {excel_row} bị trùng Tên '{clean_name}'.")
            if text_key(clean_code) in codes:
                raise InputError(f"File '{path}', dòng {excel_row} bị trùng Mã '{clean_code}'.")
            names.add(text_key(clean_name))
            codes.add(text_key(clean_code))
            devs.append(
                Dev(
                    clean_name,
                    clean_code,
                    level,
                    active_key == "y",
                    _leave_days(path, excel_row, _cell(cells, columns, "ngày nghỉ")),
                )
            )
        if not devs:
            raise InputError(f"File '{path}' không có dòng dev nào.")
        return tuple(devs)
    finally:
        workbook.close()


def _cell(cells: tuple[object, ...], columns: dict[str, int], key: str) -> object:
    index = columns.get(key)
    return cells[index] if index is not None and index < len(cells) else None


def _leave_days(path: Path, excel_row: int, value: object) -> tuple[date, ...]:
    values = (value,) if isinstance(value, (date, datetime)) else clean(value).split(",")
    days: list[date] = []
    for item in values:
        if not clean(item):
            continue
        try:
            parsed = parse_date(item)
        except ValueError as error:
            raise InputError(f"File '{path}', dòng {excel_row}, cột 'Ngày nghỉ': {error}") from error
        if parsed is not None:
            days.append(parsed)
    return tuple(days)
