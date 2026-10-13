"""Write the assignment report (assignment spec §10, bot spec §6): Tóm tắt, Phân công, Rủi ro, Thay đổi, Cảnh báo,
Tải dev."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from scrumbot.actions.workbook import date_format, finish, save, slug, table
from scrumbot.model import Level, Result
from scrumbot.model_risk import LEVEL_TEXT
from scrumbot.understanding.sizing import ORDER_TYPES

EXTRA_HEADERS = (
    "Độ lớn",
    "Giờ ước tính",
    "Lý do độ lớn",
    "Loại đơn (AI)",
    "Hạng mục (AI)",
    "Rủi ro kỹ thuật (AI)",
    "PIC",
    "Mã PIC",
    "Nguồn",
    "Lý do",
    "Dự phòng 1",
    "Dự phòng 2",
    "Rủi ro",
)
RISK_HEADERS = ("Mức", "Mã", "Nhóm", "Vị trí", "Vấn đề", "Ảnh hưởng", "Đề xuất", "Ai xử lý", "Hạn xử lý")
CHANGE_HEADERS = ("Loại", "Đơn", "Trước", "Sau")
FINDING_HEADERS = ("Mức", "Mã", "Vị trí", "Nội dung")
DEV_HEADERS = (
    "Mã",
    "Tên",
    "Level",
    "Việc dở sprint trước",
    "Ngày nghỉ",
    "Đơn giao sẵn",
    "Đơn tool đề xuất",
    "Tổng đơn",
    "Tổng giờ ước tính",
    "Danh sách đơn",
    "App đã làm",
    "Sức chứa còn lại (h)",
    "Giờ việc còn lại (h)",
)
LEVEL_NAMES = {
    Level.JUNIOR: "junior",
    Level.JUNIOR_PLUS: "junior+",
    Level.MIDDLE: "middle",
    Level.MIDDLE_PLUS: "middle+",
    Level.SENIOR: "senior",
}


def report_name(sheet_name: str, now: datetime) -> str:
    """`Sprint 10` at 2026-10-06 14:30:00 → `assignment_sprint-10_20261006-143000.xlsx` (best practices §5)."""
    return f"assignment_{slug(sheet_name)}_{now:%Y%m%d-%H%M%S}.xlsx"


def write_report(result: Result, headers: tuple[str, ...], out_dir: Path, now: datetime) -> Path:
    """The six sheets of bot spec §6, written to `<name>.tmp` and renamed (best practices §6)."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tóm tắt"
    table(summary, ("Mục", "Nội dung"), [list(row) for row in result.summary])
    _assignments_sheet(workbook.create_sheet("Phân công"), result, headers)
    _risks_sheet(workbook.create_sheet("Rủi ro"), result)
    _changes_sheet(workbook.create_sheet("Thay đổi"), result)
    table(
        workbook.create_sheet("Cảnh báo"),
        FINDING_HEADERS,
        [[f.severity.value, f.code, f.where, f.message] for f in result.findings],
        "Không có cảnh báo",
    )
    _devs_sheet(workbook.create_sheet("Tải dev"), result)
    return save(workbook, out_dir / report_name(result.sheet_name, now))


def _assignments_sheet(sheet: Worksheet, result: Result, headers: tuple[str, ...]) -> None:
    risk_codes: dict[int, list[str]] = {}
    for risk in result.risks:
        if risk.stt is not None and risk.code not in risk_codes.setdefault(risk.stt, []):
            risk_codes[risk.stt].append(risk.code)
    sheet.append([*headers, *EXTRA_HEADERS])
    for item in result.assignments:
        cells = list(item.order.row.cells[: len(headers)])
        cells += [None] * (len(headers) - len(cells))
        size = item.size
        backups = (*item.backups, "", "")
        sheet.append(
            [
                *cells,
                size.size.name + (" (mặc định)" if size.is_default else ""),
                size.size.value,
                size.reason,
                ORDER_TYPES.get(size.kind, ""),
                "; ".join(size.items),
                "; ".join(size.tech_risks),
                ", ".join(item.pic_names),
                ", ".join(item.pic_codes),
                item.source.value,
                item.reason,
                backups[0],
                backups[1],
                ", ".join(risk_codes.get(item.order.row.stt, ())),
            ]
        )
    finish(sheet)


def _risks_sheet(sheet: Worksheet, result: Result) -> None:
    rows: list[list[object]] = [
        [
            LEVEL_TEXT[risk.level],
            risk.code,
            risk.group.value,
            risk.where,
            risk.problem,
            risk.impact,
            risk.action,
            risk.owner,
            risk.due,
        ]
        for risk in result.risks
    ]
    table(sheet, RISK_HEADERS, rows, "Không có rủi ro")
    date_format(sheet, 9, 9)


def _changes_sheet(sheet: Worksheet, result: Result) -> None:
    sheet.append(list(CHANGE_HEADERS))
    if result.change_note:
        sheet.append([result.change_note])
    for change in result.changes:
        sheet.append([change.kind.value, change.order, change.before, change.after])
    if not result.changes and result.change_note.startswith("So với"):
        sheet.append(["Không có thay đổi"])
    finish(sheet)


def _devs_sheet(sheet: Worksheet, result: Result) -> None:
    rows: list[list[object]] = []
    for item in result.devs:
        orders = "Nghỉ cả sprint" if item.on_leave_all_sprint else "\n".join(item.orders)
        rows.append(
            [
                item.dev.code,
                item.dev.name,
                LEVEL_NAMES[item.dev.level],
                item.carry_over,
                item.leave_days,
                item.preassigned,
                item.proposed,
                item.preassigned + item.proposed,
                item.hours,
                orders,
                ", ".join(item.known_apps),
                item.capacity,
                item.open_hours if item.capacity is not None else None,
            ]
        )
    table(sheet, DEV_HEADERS, rows)
