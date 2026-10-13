"""Shared workbook writing: header style, widths, date format, atomic save (best practices §6)."""

from __future__ import annotations

import os
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_WIDTH = 60


def slug(name: str) -> str:
    """`Sprint 10 (10/6 - 10/17)` → `sprint-10`; only `[a-z0-9._-]` in file names (best practices §5)."""
    short = re.sub(r"\s*\(.*$", "", name.strip()) or name.strip()
    return re.sub(r"[^a-z0-9._-]", "_", re.sub(r"\s+", "-", short.lower()))


def save(workbook: Workbook, path: Path) -> Path:
    """Write `<name>.tmp`, then rename: a crash never leaves a half report that looks complete."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    workbook.save(temporary)
    os.replace(temporary, path)
    return path


def table(sheet: Worksheet, headers: tuple[str, ...], rows: list[list[object]], empty: str | None = None) -> None:
    sheet.append(list(headers))
    if not rows and empty is not None:
        sheet.append([empty])
    for row in rows:
        sheet.append(row)
    finish(sheet)


def date_format(sheet: Worksheet, first: int, last: int) -> None:
    for row in sheet.iter_rows(min_row=2, min_col=first, max_col=last):
        for cell in row:
            cell.number_format = "DD/MM/YYYY"


def finish(sheet: Worksheet) -> None:
    """Bold, frozen header; wrapped text; widths from the longest line in each column, capped."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).splitlines() or [""])
                widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), MAX_WIDTH)
