"""Write the replay workbook (bot spec §8.3) and the end-of-sprint report (bot spec §9)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from scrumbot.actions.workbook import date_format, save, slug, table
from scrumbot.model_risk import LEVEL_TEXT, CodeScore, Outcome, ReplaySprint, RiskSprint
from scrumbot.rules.replay import percent

ASSIGN_SPRINT_HEADERS = ("Sprint", "Bắt đầu", "Kết thúc", "Pool", "Đơn so sánh", "Khớp", "Tỉ lệ khớp", "Đơn mang sang")
ASSIGN_ORDER_HEADERS = (
    "Sprint",
    "Ngày tạo",
    "ID",
    "App",
    "Đơn",
    "Ưu tiên",
    "TM giao",
    "Bot giao",
    "Khớp",
    "Dự phòng 1",
    "Dự phòng 2",
    "Lý do của bot",
    "Luật TM",
)
CODE_HEADERS = (
    "Mã",
    "Đơn được báo",
    "Báo đúng",
    "Tỉ lệ báo đúng",
    "Đơn kết cục xấu",
    "Được báo trước",
    "Tỉ lệ báo đủ",
    "Báo trước (ngày, trung vị)",
)
RISK_SPRINT_HEADERS = (
    "Sprint",
    "Đơn",
    "Kết cục xấu",
    "Được báo trước",
    "Đơn được báo",
    "Báo đúng",
    "Có lịch sử trạng thái",
)
OUTCOME_HEADERS = ("Sprint", "ID", "App", "Đơn", "TM giao", "Kết cục", "Lý do", "Rủi ro đã báo")
REGISTER_HEADERS = ("Sprint", "Mốc", "Mức", "Mã", "Vị trí", "Vấn đề")
PROPOSAL_HEADERS = ("Đơn", "Bot đề xuất", "TM giao", "Khớp")


def write_replay(
    assignments: tuple[ReplaySprint, ...],
    risks: tuple[RiskSprint, ...],
    scores: tuple[CodeScore, ...],
    with_history: tuple[CodeScore, ...],
    out_dir: Path,
    now: datetime,
) -> Path:
    workbook = Workbook()
    first = workbook.active
    first.title = "Giao người theo sprint"
    rows: list[list[object]] = []
    for sprint in assignments:
        matched = sum(item.matched for item in sprint.orders)
        rows.append(
            [
                sprint.label.name,
                sprint.label.start,
                sprint.label.end,
                ", ".join(sprint.pool),
                *_rate(matched, len(sprint.orders)),
                sprint.kept,
            ]
        )
    total = sum(len(sprint.orders) for sprint in assignments)
    matched = sum(item.matched for sprint in assignments for item in sprint.orders)
    rows.append(["Tổng", None, None, "", *_rate(matched, total), sum(sprint.kept for sprint in assignments)])
    table(first, ASSIGN_SPRINT_HEADERS, rows)
    date_format(first, 2, 3)
    detail = workbook.create_sheet("Giao người theo đơn")
    detail_rows: list[list[object]] = []
    for sprint in assignments:
        for item in sprint.orders:
            row, choice = item.order.row, item.assignment
            backups = (*choice.backups, "", "")
            detail_rows.append(
                [
                    sprint.label.name,
                    row.created,
                    row.id,
                    row.app,
                    row.title,
                    row.priority_raw,
                    ", ".join(item.tm_codes),
                    choice.pic_codes[0],
                    "✓" if item.matched else "✗",
                    backups[0],
                    backups[1],
                    choice.reason,
                    ", ".join(f"dòng {line}" for line in choice.rule_lines),
                ]
            )
    table(detail, ASSIGN_ORDER_HEADERS, detail_rows)
    date_format(detail, 2, 2)
    _risk_sheets(workbook, risks, scores, with_history=with_history)
    return save(workbook, out_dir / f"replay_{now:%Y%m%d-%H%M%S}.xlsx")


def write_sprint_end(
    sprint: RiskSprint, scores: tuple[CodeScore, ...], proposals: tuple[Outcome, ...], out_dir: Path, now: datetime
) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tóm tắt"
    overall = scores[-1]
    bad = sum(item.bad for item in sprint.outcomes)
    rows = [
        ["Sprint", f"{sprint.label.name} ({sprint.label.start:%d/%m/%Y} – {sprint.label.end:%d/%m/%Y})"],
        ["Đơn", len(sprint.outcomes)],
        ["Kết cục xấu", bad],
        ["Được báo trước", overall.bad_flagged],
        ["Tỉ lệ báo đủ", _percent_text(overall.bad_flagged, overall.bad)],
        ["Đơn được báo", overall.flagged],
        ["Báo đúng", overall.flagged_bad],
        ["Tỉ lệ báo đúng", _percent_text(overall.flagged_bad, overall.flagged)],
        ["Có lịch sử trạng thái", "Có" if sprint.has_history else "Không"],
    ]
    table(summary, ("Mục", "Nội dung"), rows)
    _risk_sheets(workbook, (sprint,), scores, per_sprint=False)
    table(
        workbook.create_sheet("Đề xuất của bot"),
        PROPOSAL_HEADERS,
        [[item.title, item.proposed, ", ".join(item.actual), "✓" if item.matched else "✗"] for item in proposals],
        "Không có đề xuất nào được TM giao trong sprint này",
    )
    return save(workbook, out_dir / f"cuoi-sprint_{slug(sprint.label.name)}_{now:%Y%m%d-%H%M%S}.xlsx")


def _risk_sheets(
    workbook: Workbook,
    risks: tuple[RiskSprint, ...],
    scores: tuple[CodeScore, ...],
    *,
    per_sprint: bool = True,
    with_history: tuple[CodeScore, ...] | None = None,
) -> None:
    table(workbook.create_sheet("Rủi ro theo loại"), CODE_HEADERS, _score_rows(scores))
    if with_history is not None:
        table(workbook.create_sheet("Rủi ro theo loại (có lịch sử)"), CODE_HEADERS, _score_rows(with_history))
    if per_sprint:
        rows: list[list[object]] = []
        for sprint in risks:
            flagged = [item for item in sprint.outcomes if item.flagged]
            rows.append(
                [
                    sprint.label.name,
                    len(sprint.outcomes),
                    sum(item.bad for item in sprint.outcomes),
                    sum(item.bad for item in flagged),
                    len(flagged),
                    sum(item.bad for item in flagged),
                    "Có" if sprint.has_history else "Không",
                ]
            )
        table(workbook.create_sheet("Rủi ro theo sprint"), RISK_SPRINT_HEADERS, rows)
    outcome_rows: list[list[object]] = []
    register_rows: list[list[object]] = []
    for sprint in risks:
        for item in sprint.outcomes:
            row = item.order.row
            outcome_rows.append(
                [
                    sprint.label.name,
                    row.id,
                    row.app,
                    row.title,
                    ", ".join(item.tm_codes),
                    "Xấu" if item.bad else "Tốt",
                    ", ".join(item.reasons),
                    ", ".join(f"{code} ({day:%d/%m})" for code, day in item.flagged),
                ]
            )
        for point in sprint.checkpoints:
            for risk in point.risks:
                register_rows.append(
                    [
                        sprint.label.name,
                        f"{point.at:%d/%m/%Y %H:%M}",
                        LEVEL_TEXT[risk.level],
                        risk.code,
                        risk.where,
                        risk.problem,
                    ]
                )
    table(workbook.create_sheet("Kết cục đơn"), OUTCOME_HEADERS, outcome_rows)
    table(workbook.create_sheet("Sổ rủi ro các mốc"), REGISTER_HEADERS, register_rows, "Không có rủi ro")


def _score_rows(scores: tuple[CodeScore, ...]) -> list[list[object]]:
    return [
        [
            score.code,
            score.flagged,
            score.flagged_bad,
            _percent_text(score.flagged_bad, score.flagged),
            score.bad,
            score.bad_flagged,
            _percent_text(score.bad_flagged, score.bad),
            score.lead_median,
        ]
        for score in scores
    ]


def _rate(matched: int, total: int) -> list[object]:
    return [total, matched, f"{percent(matched, total)}%" if total else ""]


def _percent_text(part: int, whole: int) -> str:
    return f"{percent(part, whole)}%" if whole else ""
